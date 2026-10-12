"""Tests for Talk2Data Phase 2 Charts, Tables, Disclosure Controls, and Secure XLSX/CSV Exports."""
import io
import openpyxl
import pytest
from decimal import Decimal

from talk2data.domain.facts import ProtectedFact
from talk2data.domain.models import ApiActor, Turn, TurnStatus
from talk2data.export.export import SecureExportService, sanitize_formula_injection
from talk2data.presentation.charts import ChartGenerator
from talk2data.presentation.fallback import RenderedAnswerView, RenderedKpiCard
from talk2data.security.auth import AuthorizationError
from talk2data.storage.repository import PersistentApplicationStore


def test_formula_injection_sanitization():
    """Verify dangerous spreadsheet formula prefixes are safely escaped with single quote."""
    assert sanitize_formula_injection("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_formula_injection("+123456") == "'+123456"
    assert sanitize_formula_injection("-CMD|' /C calc'!A0") == "'-CMD|' /C calc'!A0"
    assert sanitize_formula_injection("@SUM(B1)") == "'@SUM(B1)"
    assert sanitize_formula_injection("\tcmd.exe") == "'\tcmd.exe"
    # Safe text left unchanged
    assert sanitize_formula_injection("Electronics") == "Electronics"
    assert sanitize_formula_injection(12345) == 12345


def test_chart_generation_with_disclosure_suppression():
    """Verify that generated charts reflect protected facts and apply suppression controls to tooltips and labels."""
    generator = ChartGenerator()

    facts = [
        ProtectedFact(
            fact_id="f1",
            metric="sales_revenue",
            calculation_type="GROUPED",
            value=Decimal("50000.00"),
            unit="USD",
            period="2024",
            dimension_values={"category": "Electronics"},
            is_suppressed=False,
        ),
        ProtectedFact(
            fact_id="f2",
            metric="sales_revenue",
            calculation_type="GROUPED",
            value=Decimal("0.00"),
            unit="USD",
            period="2024",
            dimension_values={"category": "RareGadgets"},
            is_suppressed=True,
            suppression_reason="Cell order count below k=5 privacy threshold",
        ),
    ]

    charts = generator.generate_charts(facts)
    assert len(charts) == 1
    bar_chart = charts[0]
    assert bar_chart["type"] == "bar"
    assert bar_chart["has_suppression"] is True
    assert "RareGadgets [SUPPRESSED]" in bar_chart["labels"]
    assert bar_chart["values"][1] is None
    assert "[SUPPRESSED]" in bar_chart["tooltips"][1]
    assert "privacy threshold" in bar_chart["tooltips"][1]


def test_secure_xlsx_export_and_formula_injection(test_setup):
    """Verify XLSX export generates valid workbook, escaping formula injection and preserving suppression."""
    store = PersistentApplicationStore(":memory:")
    security = test_setup["security"]
    actor = test_setup["actor_a"]

    conv_id = store.create_conversation(actor, title="Export Test")
    turn_id = "turn_export_01"

    t = Turn.create(
        conversation_id=conv_id,
        actor=actor,
        question="=HYPERLINK('http://malicious.com','Click Me')",
        source_id="finance_dw",
        turn_id=turn_id,
    )
    store.save_turn(t, actor)

    view = RenderedAnswerView(
        summary_text="=cmd|'/C calc'!A0",
        kpi_cards=[RenderedKpiCard(title="Revenue", value="+9999", unit="USD", period="2024")],
        table_columns=["Category", "=FormulaCol", "Revenue"],
        table_rows=[
            {"Category": "Electronics", "=FormulaCol": "@DANGEROUS()", "Revenue": "$50,000.00"},
            {"Category": "-SuppressedItem", "=FormulaCol": "Safe", "Revenue": "[SUPPRESSED]"},
        ],
        assumptions=["=MaliciousAssumption()"],
        limitations=["Standard limitation"],
        is_suppressed=True,
        evidence_digest="digest_export_test",
    )

    store.publish_answer_transactional(turn_id=turn_id, actor=actor, view=view, version=1)

    store.security = security
    store.catalogs = test_setup["catalogs"]
    import json
    catalog = store.catalogs.get_catalog("finance_dw")
    store.conn.execute("UPDATE turns SET pinned_release_id=?, execution_state_json=? WHERE turn_id=?",
        (catalog.release_id, json.dumps({"intent": {"measure": "sales_revenue", "dimensions": [], "filters": []}}), turn_id))
    store.conn.commit()
    export_service = SecureExportService(store=store, security=security, catalogs=store.catalogs)
    xlsx_bytes = export_service.export_xlsx(actor, turn_id)

    # Inspect exported workbook
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Analysis Results" in wb.sheetnames
    assert "Governance & Scope" in wb.sheetnames

    ws_data = wb["Analysis Results"]
    # Check that headers had formula prefixes sanitized
    assert ws_data.cell(row=1, column=2).value == "'=FormulaCol"

    # Check that cells had formula prefixes sanitized
    assert ws_data.cell(row=2, column=2).value == "'@DANGEROUS()"
    assert ws_data.cell(row=3, column=1).value == "'-SuppressedItem"
    assert ws_data.cell(row=3, column=3).value == "[SUPPRESSED]"

    # Governance sheet
    ws_gov = wb["Governance & Scope"]
    assert ws_gov.cell(row=2, column=2).value == "'=HYPERLINK('http://malicious.com','Click Me')"
    assert ws_gov.cell(row=3, column=2).value == "''=cmd|'/C calc'!A0" or ws_gov.cell(row=3, column=2).value.startswith("'")


def test_dynamic_permission_revocation_blocks_export(test_setup):
    """Verify dynamic permission revocation at export time blocks both XLSX and CSV exports with AuthorizationError."""
    store = PersistentApplicationStore(":memory:")
    security = test_setup["security"]
    actor = test_setup["actor_a"]

    conv_id = store.create_conversation(actor, title="Revocation Export Test")
    turn_id = "turn_rev_01"

    t = Turn.create(conversation_id=conv_id, actor=actor, question="Revenue?", source_id="finance_dw", turn_id=turn_id)
    store.save_turn(t, actor)

    view = RenderedAnswerView(
        summary_text="Total revenue was $80,000.00.",
        kpi_cards=[RenderedKpiCard(title="Revenue", value="$80,000.00", unit="USD", period="2024")],
        table_columns=["Metric", "Value"],
        table_rows=[{"Metric": "Revenue", "Value": "$80,000.00"}],
        assumptions=[],
        limitations=[],
        is_suppressed=False,
        evidence_digest="digest_rev_test",
    )
    store.publish_answer_transactional(turn_id=turn_id, actor=actor, view=view, version=1)

    store.security = security
    store.catalogs = test_setup["catalogs"]
    import json
    catalog = store.catalogs.get_catalog("finance_dw")
    store.conn.execute("UPDATE turns SET pinned_release_id=?, execution_state_json=? WHERE turn_id=?",
        (catalog.release_id, json.dumps({"intent": {"measure": "sales_revenue", "dimensions": [], "filters": []}}), turn_id))
    store.conn.commit()
    export_service = SecureExportService(store=store, security=security, catalogs=store.catalogs)

    # Initial export works
    csv_data = export_service.export_csv(actor, turn_id)
    assert "Revenue" in csv_data

    # Dynamically revoke tenant's access to the data source
    security.revoke_tenant_source(actor.tenant_id, "finance_dw")

    # Subsequent export attempts must fail immediately
    with pytest.raises(AuthorizationError, match="Source access denied"):
        export_service.export_xlsx(actor, turn_id)

    with pytest.raises(AuthorizationError, match="Source access denied"):
        export_service.export_csv(actor, turn_id)
