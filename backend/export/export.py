"""Secure export service generating XLSX and CSV files with formula injection protection and disclosure controls."""
from __future__ import annotations
import csv
import io
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill

from backend.domain.models import ApiActor
from backend.security.auth import AuthorizationError, SecurityRuntime
from backend.storage.repository import PersistentApplicationStore


def sanitize_formula_injection(val: Any) -> Any:
    """Prepend single quote if a string begins with dangerous spreadsheet formula triggers."""
    if isinstance(val, str):
        if val.startswith(("=", "+", "-", "@", "\t", "\r")):
            return f"'{val}"
    return val


class SecureExportService:
    """Exports approved analysis datasets to XLSX and CSV while enforcing privacy and security controls."""

    def __init__(
        self,
        store: PersistentApplicationStore,
        security: SecurityRuntime,
        catalogs: Optional[Any] = None,
    ):
        self.store = store
        self.security = security
        self.catalogs = catalogs

    def _verify_export_permissions(self, actor: ApiActor, turn: dict[str, Any]) -> None:
        """Verify current permissions for source and catalog assets before exporting."""
        self.store.verify_answer_permissions(actor, turn)
        if self.catalogs:
            cat = self.catalogs.get_catalog(turn["source_id"], turn["pinned_release_id"])
            if cat.disclosure_policy and not cat.disclosure_policy.get("export_allowed", True):
                raise AuthorizationError("Dataset policy prohibits exports")

    def export_xlsx(self, actor: ApiActor, turn_id: str) -> bytes:
        """Generate XLSX workbook from published answer, checking current permissions."""
        answer = self.store.get_published_answer(actor, turn_id)
        turn = self.store.get_turn(actor, turn_id)

        # Dynamic permission re-check at export time
        self._verify_export_permissions(actor, turn)

        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = "Analysis Results"

        # Header row
        columns = answer["table_columns"]
        ws_data.append([sanitize_formula_injection(col) for col in columns])

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        for col_idx in range(1, len(columns) + 1):
            cell = ws_data.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        # Rows (sanitizing formula triggers and preserving [SUPPRESSED])
        for row_dict in answer["table_rows"]:
            row_vals = []
            for col in columns:
                val = row_dict.get(col, "")
                if val == "[SUPPRESSED]":
                    row_vals.append("[SUPPRESSED]")
                else:
                    row_vals.append(sanitize_formula_injection(val))
            ws_data.append(row_vals)

        # Meta & Disclosure sheet
        ws_meta = wb.create_sheet(title="Governance & Scope")
        ws_meta.append(["Attribute", "Value"])
        ws_meta.cell(row=1, column=1).font = header_font
        ws_meta.cell(row=1, column=1).fill = header_fill
        ws_meta.cell(row=1, column=2).font = header_font
        ws_meta.cell(row=1, column=2).fill = header_fill

        ws_meta.append(["Question", sanitize_formula_injection(turn["question_text"])])
        ws_meta.append(["Summary", sanitize_formula_injection(answer["summary_text"])])
        ws_meta.append(["Evidence Digest", sanitize_formula_injection(answer["evidence_digest"])])
        ws_meta.append(["Suppression Applied", "YES" if answer["is_suppressed"] else "NO"])

        for idx, asm in enumerate(answer["assumptions"], 1):
            ws_meta.append([f"Assumption {idx}", sanitize_formula_injection(asm)])
        for idx, lim in enumerate(answer["limitations"], 1):
            ws_meta.append([f"Limitation {idx}", sanitize_formula_injection(lim)])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def export_csv(self, actor: ApiActor, turn_id: str) -> str:
        """Generate CSV string from published answer with formula injection sanitization."""
        answer = self.store.get_published_answer(actor, turn_id)
        turn = self.store.get_turn(actor, turn_id)

        # Dynamic permission re-check
        self._verify_export_permissions(actor, turn)

        output = io.StringIO()
        writer = csv.writer(output)

        columns = answer["table_columns"]
        writer.writerow([sanitize_formula_injection(c) for c in columns])

        for row_dict in answer["table_rows"]:
            writer.writerow([sanitize_formula_injection(row_dict.get(c, "")) for c in columns])

        return output.getvalue()
